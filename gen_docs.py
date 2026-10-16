# -*- coding: utf-8 -*-
"""
生成两份 Word 文档：
  1. 第四讲-讲解.docx       —— 课件逐节讲解 + 术语解释
  2. Agent功能说明.docx      —— 本仓库记忆机器人 Agent 的主要功能说明
输出到桌面与本仓库 docs/ 各一份。
运行：python gen_docs.py
"""
import os
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DESKTOP = os.path.join(os.path.expanduser("~"), "Desktop")
DOCS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "docs")


# ---------- 排版辅助 ----------

def set_cn_fonts(doc):
    """让中文用微软雅黑渲染，避免默认字体发虚。"""
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Title"):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = "Calibri"
        rpr = st.element.get_or_add_rPr()
        rpr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    doc.styles["Normal"].font.size = Pt(11)


def h(doc, text, level=1):
    doc.add_heading(text, level)


def p(doc, text, bold=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    return para


def bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def code(doc, text):
    for line in text.strip("\n").split("\n"):
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.left_indent = Pt(18)
        run = para.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x1F, 0x3F, 0x7F)


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, head in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = head
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    return t


def save(doc, filename):
    os.makedirs(DOCS_DIR, exist_ok=True)
    local = os.path.join(DOCS_DIR, filename)
    doc.save(local)
    shutil.copy(local, os.path.join(DESKTOP, filename))
    print("已生成:", local, "和桌面副本")


# ================= 文档一：第四讲-讲解.docx =================

def build_lecture_doc():
    doc = Document()
    set_cn_fonts(doc)

    title = doc.add_heading("《Agent 的记忆管理 · 第四讲》讲解", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p(doc, "课程：中山大学《软工新生研讨实践》（李煜政、郑沛霖，2026.10）　　整理日期：2026-10-16")
    p(doc, "配套代码仓库：agent-lecture4（memory_bot.py / llm.py / gen_docs.py）")

    # ---------- 一 ----------
    h(doc, "一、本讲主线：大模型没有记忆，记忆是调用方做出来的", 1)
    p(doc, "第三讲解决了“机器人怎么收发消息”，这一讲解决“机器人怎么记得你说过话”。"
           "课件用一个实验开场：先告诉模型“我叫小明”，紧接着问“我叫什么名字”——模型答不上来。"
           "原因是大模型接口是无状态（stateless）的：每次推理只基于这一次请求，"
           "请求与请求之间互不相干，模型本身不会记住任何上一轮的内容。")
    p(doc, "所以“记忆”不是模型的能力，而是调用方（我们写的程序）的功课："
           "把历史对话想办法塞进每一次新请求里。整节课就是围绕“怎么塞、塞多少、塞哪些”"
           "一步步演进，最后得到一张五阶段总结表：")
    table(doc,
          ["阶段", "做法", "留下的问题"],
          [["1", "完整历史：每轮把全部对话重新发一遍", "对话变长后达到上下文窗口上限"],
           ["2", "滑动窗口：只带最近 N 轮", "丢失更早的内容"],
           ["3", "压缩摘要：太长就先让模型总结成一段话", "丢失具体细节"],
           ["4", "持久化存储 + 关键词选取：写入文件，按需读取", "关键词匹配程度有限（只认字面）"],
           ["5", "延伸：向量检索 / Agentic Search", "需要向量数据库 / 需要工具调用能力"]])

    # ---------- 二 ----------
    h(doc, "二、上下文与基本调用：messages 与 role", 1)

    h(doc, "2.1 无状态实验", 2)
    p(doc, "连续两次调用 API：第一次说“我叫小明”，第二次问“我叫什么名字”。"
           "第二次请求里没有任何关于“小明”的信息，模型自然回答不知道。"
           "结论：想让模型“记得”，必须在每次请求里把相关历史一起发过去。")

    h(doc, "2.2 messages 数组与 role 试错", 2)
    p(doc, "发给模型的对话体是一个 messages 数组，每条消息都要带 role（角色）。课件的试错过程：")
    bullets(doc, [
        "不写 role → 接口直接报 400；",
        "自创 role（如 me）→ 同样 400，role 只允许固定的几个值；",
        "全部标成 user → 能跑，但模型分不清哪句是谁说的，效果差。",
    ])
    p(doc, "三个常用角色：system（系统设定/人设）、user（用户输入）、assistant（模型回答）。")

    h(doc, "2.3 role 的底层原理：ChatML", 2)
    p(doc, "role 不是玄学。课件展示了 token 层面的展开：messages 数组在送进模型前会被拼成一串文本，"
           "每条消息用特殊 token 包裹（如 151644 是开始标记、872 代表 user、77091 代表 assistant、"
           "198 是换行、151645 是结束标记），最后留一个 assistant 开始标记，"
           "模型就从这里往后自回归预测——逐个 token 地生成回答。"
           "这套格式叫 ChatML，role 的作用就是让模型在训练好的“剧本格式”里分清谁在说话。")

    # ---------- 三 ----------
    h(doc, "三、阶段一：携带全部历史", 1)
    p(doc, "最直白的记忆方案：程序里维护一个 messages 列表，每轮对话把 user 和 assistant 两条 append 进去，"
           "下次请求整个列表全量重发：")
    code(doc, '''
messages = []
while True:
    text = input("你: ")
    messages.append({"role": "user", "content": text})
    reply = ask(messages)
    messages.append({"role": "assistant", "content": reply})''')
    p(doc, "这样模型每轮都能看到从第一句话开始的完整对话，“记忆”就有了。代价在下一节。")

    h(doc, "3.1 Token 计费", 2)
    bullets(doc, [
        "API 按 token 计费，一个汉字大约 1~2 个 token；",
        "每次请求按完整 messages 的 token 数计费——带的历史越长，单次越贵；",
        "上下文窗口 = 单次请求允许的 token 上限，撞顶就报错；",
        "打印返回包里的 usage 字段，可以直观看到每次请求的 token 消耗。",
    ])

    h(doc, "3.2 缓存命中（前缀缓存）", 2)
    p(doc, "每次发送的历史前缀和上一次完全相同，服务商后台会把相同部分的计算缓存起来，"
           "命中缓存的 token 单价更低。方舟等平台的 usage 里有 cached_tokens 字段，"
           "DeepSeek 则给出 prompt_cache_hit_tokens / prompt_cache_miss_tokens。"
           "两个推论：system 人设要固定放在最前面；历史只能往后追加、不能改动中间——"
           "前缀一旦变了，后面全部按原价重新计费。")

    h(doc, "3.3 撞顶问题", 2)
    p(doc, "对话越长，列表越长，单次成本越高，最终触碰上下文窗口上限。"
           "早期 ChatGPT 窗口小，长对话聊着聊着就“忘了开头”，就是这个原因。"
           "于是引出后面三种“省着用”的记忆方案。")

    # ---------- 四 ----------
    h(doc, "四、阶段二：滑动窗口", 1)
    p(doc, "只带最近 N 轮，更早的直接丢弃：")
    code(doc, '''
N = 3
reply = ask(messages[-2 * N:] + [{"role": "user", "content": text}])''')
    p(doc, "一轮 = user + assistant 两条，所以是 -2*N。请求长度从此可控，"
           "代价是早期对话信息永久丢失——问了 N 轮之前的事，模型就又不知道了。")

    # ---------- 五 ----------
    h(doc, "五、阶段三：压缩摘要", 1)
    p(doc, "历史超过阈值时，先让模型自己把前面的对话总结成一段话，再用摘要替换原文：")
    code(doc, '''
LIMIT = 6
if len(messages) > LIMIT:
    req = messages + [{"role": "user", "content": "把以上对话总结成一段话"}]
    summary = ask(req)
    messages = [{"role": "system", "content": "之前的对话摘要: " + summary}]
    print("[已压缩成摘要]")''')
    p(doc, "主要内容保住了，具体细节（原话、数字、名单）会在摘要中丢失。"
           "产品实践：Claude Code 在上下文接近上限时会自动压缩；"
           "ChatGPT 也会对长对话的前文生成摘要。另一个思路是先把历史持久化，需要时再读取——即下一阶段。")

    # ---------- 六 ----------
    h(doc, "六、阶段四：持久化存储与选取", 1)

    h(doc, "6.1 按天写文件", 2)
    p(doc, "每天一个文本文件，一轮对话追加两行，程序重启后从当天文件读回，记忆还在：")
    code(doc, '''
os.makedirs("memory", exist_ok=True)
path = "memory/" + datetime.date.today().isoformat() + ".txt"
# 读回：line.rstrip("\\n").split(": ", 1) 还原成 messages
# 追加："user: " + text + "\\n" + "assistant: " + reply + "\\n"''')

    h(doc, "6.2 全部读回又会撞顶 → 有选择地读取", 2)
    p(doc, "日子久了文件越积越多，全部读回等于重新面对窗口上限。"
           "选取思路：从新问题里提取关键词 → 匹配历史文件 → 只有命中的文件才带进请求。"
           "而提取关键词的前提是先分词。")

    h(doc, "6.3 jieba 结巴分词", 2)
    bullets(doc, [
        "中文没有空格，得先切词。jieba 是最常用的中文分词库（pip install jieba）；",
        "原理一：基于词典，从多种切法里选最合理的一种；",
        "原理二：词典未收录的词，按统计规律推测；",
        "例子：jieba.cut(\"南京市长江大桥\") → ['南京市', '长江大桥']。",
    ])

    h(doc, "6.4 关键词选取与拼装", 2)
    code(doc, '''
keys = [w for w in jieba.cut(question) if len(w) > 1]   # 去掉单字和标点
picked = []
for path in sorted(glob.glob("memory/*.txt")):
    text = open(path, encoding="utf-8").read()
    if any(k in text for k in keys):
        picked.append(text)          # 命中任何一个词，整天内容带回
content = "\\n".join(picked) + "\\n以上是历史记录，请回答：" + question''')
    p(doc, "命中的日期文件按日期顺序拼接，末尾加一句引导语，再附上本次问题，"
           "整条作为一条 user 消息发出。注意：模型无法区分各段分别来自哪一天。")

    h(doc, "6.5 关键词匹配的局限", 2)
    p(doc, "课件反例：用户问“这东西坏了能换吗”，提取出的关键词是“东西”，"
           "而文件里记的是“支持七天无理由退货”——字面不重合，命中失败。"
           "关键词匹配只识别字面，不识别语义，这是本阶段留下的问题。")

    h(doc, "6.6 索引目录与规约文件", 2)
    bullets(doc, [
        "Claude Code 的做法：项目根目录放 CLAUDE.md，启动时自动加载；",
        "经验拆成独立文件时，另用一个 index.md 作目录，一行一条（如 - [2025-10-05.txt] 订单号和快递），"
        "启动只载入目录，命中后再去读对应文件——比全文盲扫省得多；",
        "各家的规约文件（灵魂文件）原本互不兼容：CLAUDE.md / .cursor/rules / GEMINI.md；",
        "AGENTS.md 已成为开放标准，一份通用；Claude Code 一度不读被用户控诉，2026 年 9 月起也已兼容。",
    ])

    h(doc, "6.7 全文检索 grep", 2)
    code(doc, '''
grep -r "退款" memory/     # 递归搜目录
grep -rn "退款" memory/    # 连行号一起显示''')
    p(doc, "对代码符号、报错信息这类精确匹配场景非常有效。"
           "真实产品里，搜什么关键词由模型自己决定——这就是延伸部分的 Agentic Search。")

    # ---------- 七 ----------
    h(doc, "七、延伸：Agentic Search 与向量检索", 1)

    h(doc, "7.1 Agentic Search", 2)
    p(doc, "让模型自己决定搜索关键词，再调用 grep 等工具去查；查到结果后还能顺着新线索再来一轮。"
           "例如问“上次看周杰伦演唱会是什么时候”，模型自己决定搜“周杰伦”“演唱会”。"
           "前提是模型具备工具调用（Function/Tool Calling）能力。")

    h(doc, "7.2 向量检索与 RAG", 2)
    bullets(doc, [
        "向量检索：把文字转成一串数字（向量/Embedding），语义越近，数字距离越近——"
        "“退款”和“退货”很近，和“演唱会”很远，正好补上关键词只认字面的短板；",
        "RAG（检索增强生成）流程：历史切块 → 转成向量 → 存入向量数据库 → "
        "问题也转成向量 → 取距离最近的几条带进请求；",
        "可自行体验的平台：coze.cn（扣子）、Dify。",
    ])

    h(doc, "7.3 产品实践对照", 2)
    bullets(doc, [
        "Claude Code / Codex：不建向量索引，直接让模型用 grep 现搜；",
        "Cursor：曾用语义索引，后来也把 grep 列为核心工具；",
        "启示：工程上“简单直接的精确检索 + 让模型自己多搜几轮”常常胜过复杂的向量方案。",
    ])

    # ---------- 八 ----------
    h(doc, "八、课后任务（本仓库已完成）", 1)
    bullets(doc, [
        "跑一个程序：今天的 Python 代码，或接入上一讲的机器人；",
        "问它两个问题，第二个问题在第一个的答案基础上追问；",
        "示例：先给名字让它写藏头诗，再问它解读诗意；",
        "截图提交。交了就有分，重在实践，不要抄，不要卷；",
        "之前的任务接受补交，但要尽快，以免后面跟不上。",
    ])
    p(doc, "本仓库的完成方式：memory_bot.py 把五段记忆方案合进一个命令行机器人，"
           "按课件示例演示“藏头诗 → 追问诗意”，并追加演示了重启后记忆仍在（持久化）"
           "与往期记录按需命中（关键词选取），完整运行记录见 demo_run1.txt / demo_run2.txt。")

    # ---------- 九 ----------
    h(doc, "九、术语解释", 1)
    p(doc, "按课件出现顺序整理，供复习速查：")
    table(doc,
          ["术语", "解释"],
          [["无状态 / stateless", "接口不保存任何会话信息，每次请求相互独立。大模型 API 皆如此，“记忆”全靠调用方在请求里自带历史。"],
           ["API", "应用程序编程接口。本课指大模型提供的 HTTP 接口（/chat/completions），程序按约定格式发请求、收回答。"],
           ["messages", "发给大模型的对话数组，元素是 {role, content} 字典，按时间顺序排列，是“记忆”的载体。"],
           ["role / 角色", "消息的说话人身份，只允许固定取值：system（系统设定）、user（用户）、assistant（模型）。不写或自创都会 400。"],
           ["system 角色 / 人设", "system 消息是“导演给演员的剧本”，规定模型身份与口吻。缓存角度应固定放在 messages 最前面。"],
           ["ChatML", " messages 在模型内部的展开格式：每条消息用特殊 token 包裹（开始标记 + 角色 + 内容 + 结束标记），末尾留 assistant 标记让模型接着写。"],
           ["特殊 token", "控制格式的保留符号，如课件中的 151644（开始）、151645（结束）、198（换行）、872（user）、77091（assistant）。"],
           ["自回归预测", "模型逐 token 生成：每预测下一个 token，都基于前面已有的全部 token。回答就是这样一个一个字“长”出来的。"],
           ["token", "模型处理文本的最小单位，一个汉字约 1~2 个 token。既是计费单位，也是上下文长度的度量单位。"],
           ["token 计费", "API 按请求+回答的 token 总数收费；携带全部历史意味着对话越长，单次请求越贵。"],
           ["上下文窗口", "单次请求允许的 token 上限（输入+输出）。撞顶即报错，是一切记忆压缩方案的根本动因。"],
           ["usage", "API 返回包里的用量字段，含 prompt_tokens、completion_tokens 等，用来观察每次请求的消耗。"],
           ["缓存命中 / 前缀缓存", "本次请求的历史前缀与上次相同的部分，服务商直接复用上次计算结果，单价更低。前提：历史只追加、不改动。"],
           ["cached_tokens / prompt_cache_hit_tokens", "usage 里的缓存命中字段（方舟 / DeepSeek 各自的叫法），可用来验证前缀缓存是否生效。"],
           ["滑动窗口", "记忆策略二：只携带最近 N 轮（messages[-2*N:]），更早内容直接丢弃。省 token，但丢失早期信息。"],
           ["压缩摘要", "记忆策略三：历史超阈值时让模型把前文总结成一段话，用摘要替换原文。保留大意，丢失细节。"],
           ["持久化存储", "记忆策略四（上）：把对话按天写入本地文件（memory/日期.txt），程序重启后读回，记忆跨进程存活。"],
           ["关键词选取", "记忆策略四（下）：新问题提取关键词，只把命中的历史文件带进请求，避免全部读回撞顶。"],
           ["分词", "把连续的中文切成词序列，是关键词提取的前置步骤（中文没有空格）。"],
           ["jieba / 结巴分词", "最常用的 Python 中文分词库：基于词典选最合理切法 + 对未登录词按统计规律推测。"],
           ["字面匹配 vs 语义匹配", "“这东西坏了能换吗”匹配不到“七天无理由退货”——关键词只认字面重合，不认意思相近，是关键词路线的根本局限。"],
           ["索引目录 / index.md", "给记忆文件建的一行一条的目录：启动只载目录，命中后再读具体文件，省 token 也省时间。Claude Code 的 CLAUDE.md 同理。"],
           ["规约文件 / 灵魂文件", "写给 AI 看的项目说明书：CLAUDE.md、.cursor/rules、GEMINI.md 等，规定项目背景与行为准则。"],
           ["AGENTS.md", "规约文件的开放标准，一份文件各家工具通用；2026 年 9 月起 Claude Code 也已兼容。"],
           ["grep / 全文检索", "命令行文本搜索工具（grep -r 递归、-n 带行号）。对代码符号、报错信息等精确匹配场景非常有效。"],
           ["Agentic Search", "让模型自己决定搜什么关键词、调用 grep 等工具，并根据结果继续下一轮搜索。前提：模型具备工具调用能力。"],
           ["工具调用 / Tool Calling", "模型在回答中请求调用外部工具（搜索、grep、计算器……），由程序执行后把结果喂回模型的机制。"],
           ["向量 / Embedding", "把文字映射成一串数字，语义越相近的数字距离越近，让“意思对得上”的内容可以被算出来。"],
           ["向量数据库", "专门存向量并支持“按距离找最近”的数据库，是向量检索的存储底座。"],
           ["RAG / 检索增强生成", "Retrieval-Augmented Generation：历史切块 → 转向量 → 入库 → 问题转向量 → 取最近几条带进请求，让模型据以作答。"],
           ["coze / Dify", "零代码搭建 AI 应用的平台（扣子 coze.cn、Dify），课件建议用它们亲手体验 RAG 流程。"],
           ["记忆五阶段", "本课主线：完整历史 → 滑动窗口 → 压缩摘要 → 持久化+选取 →（延伸）向量检索 / Agentic Search，每一阶段都为解决上一阶段的遗留问题而生。"]])

    save(doc, "第四讲-讲解.docx")


# ================= 文档二：Agent功能说明.docx =================

def build_agent_doc():
    doc = Document()
    set_cn_fonts(doc)

    title = doc.add_heading("记忆机器人：第四讲课后任务 Agent 功能说明", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p(doc, "对应课后任务：“跑一个程序，问它两个问题，第二个问题在第一个的答案基础上追问。”　　整理日期：2026-10-16")
    p(doc, "仓库：agent-lecture4　　语言：Python 3　　依赖：requests / jieba（可选）")

    h(doc, "一、Agent 是什么", 1)
    p(doc, "本 Agent 是一个带记忆的命令行对话机器人。大模型接口本身无状态，"
           "所谓的“记忆”全部出自本程序的工程设计：它在本地维护对话历史，"
           "每轮请求按策略挑选合适的历史片段一起发给模型，并把每轮对话追加写入按天命名的本地文件，"
           "重启后读回。由此机器人能跨轮次、跨进程记住用户说过的事。")

    h(doc, "二、总体架构", 1)
    code(doc, '''
        你（键盘输入）
            │
            ▼
     memory_bot.py ──► llm.py ──► 大模型（OpenAI 兼容接口）
            │
            ├─ 会话内历史：完整历史 / 滑动窗口 / 压缩摘要（三档可选）
            ├─ 持久化：memory/今天日期.txt 逐轮追加，启动时读回
            └─ 按需选取：jieba 分词提关键词，命中往期文件才带入请求''')

    h(doc, "三、主要功能清单", 1)

    h(doc, "3.1 会话内记忆（三档策略，MEMORY_MODE 切换）", 2)
    bullets(doc, [
        "完整历史（full，默认）：每轮把 user/assistant 两条追加进 messages，下次全量重发——"
        "模型看得到从第一句话起的所有内容；",
        "滑动窗口（window）：只携带最近 N=3 轮（messages[-2*N:]），更早的丢弃，请求长度可控；",
        "压缩摘要（compress）：历史超过 LIMIT=6 条时，先让模型把前文总结成一段话，"
        "再以 system 消息“之前的对话摘要：……”替换原文，并打印 [已压缩成摘要]；",
        "token 观察：每轮打印本次携带条数、prompt/completion token 数与缓存命中数，"
        "课件里“计费”和“前缀缓存”两个概念可以肉眼验证。",
    ])

    h(doc, "3.2 持久化存储（跨重启记忆）", 2)
    bullets(doc, [
        "每轮结束立刻向 memory/YYYY-MM-DD.txt 追加两行（user / assistant），程序崩溃也不丢当轮；",
        "启动时读回当天文件还原 messages，并提示 [记忆] 从今天文件读回 N 条历史；",
        "因此第二天重新打开程序，机器人依然记得昨天的对话。",
    ])

    h(doc, "3.3 往期记忆关键词选取", 2)
    bullets(doc, [
        "提问先经 jieba 分词提取关键词（去掉单字与标点）；",
        "遍历 memory/ 下今天以外的日期文件，任一关键词命中即把该文件整天内容带入请求，"
        "并打印 [记忆] 关键词命中往期文件；",
        "带入方式：作为 system 消息“往期的历史记录”置于最前，兼顾三种会话内策略；",
        "无 jieba 自动降级：改用连续汉字的二字滑窗做关键词，效果略差但功能不缺席。",
    ])

    h(doc, "3.4 大模型接入（llm.py，与第三讲共用）", 2)
    bullets(doc, [
        "OpenAI 兼容接口：POST {LLM_BASE}/chat/completions，换服务商只改配置不改代码；",
        "配置三级回退：环境变量 LLM_KEY / LLM_BASE / LLM_MODEL 优先，"
        "其次读 llm_key.txt / llm_base.txt / llm_model.txt，最后回落默认值；",
        "故障兜底：模型报错时机器人提示错误而非沉默，主循环不死。",
    ])

    h(doc, "3.5 演示记录（课后任务证据）", 2)
    bullets(doc, [
        "demo_run1.txt：同一会话内——①“我叫王小明，请用我的名字写藏头诗”②“这首诗描写了什么意境？"
        "逐句解读诗意”（第二问完全依赖第一问的答案）③“我叫什么名字？”验证记忆；"
        "第三轮 token 日志可见缓存命中=256，前缀缓存肉眼生效；",
        "demo_run2.txt：重启程序后——机器人从当天文件读回 6 条历史，正确回答名字与藏头诗（持久化）；"
        "再问“上次看周杰伦演唱会是什么时候”，关键词命中往期文件并据以作答（选取）。",
    ])

    h(doc, "四、安全设计", 1)
    bullets(doc, [
        "密钥零入库：llm_key.txt 列入 .gitignore，永不提交；环境变量注入则磁盘不留明文；",
        "聊天数据不出本机：memory/ 目录同样列入 .gitignore，个人对话不进仓库；",
        "密钥即密码：README 与文档均显式提醒不外发、不截图外传。",
    ])

    h(doc, "五、运行方式速查", 1)
    table(doc,
          ["步骤", "命令 / 操作"],
          [["装依赖", "pip install -r requirements.txt"],
           ["配模型密钥", "set/export LLM_KEY=sk-...（或新建 llm_key.txt）"],
           ["完整历史模式（默认）", "python memory_bot.py"],
           ["滑动窗口模式", "MEMORY_MODE=window python memory_bot.py"],
           ["压缩摘要模式", "MEMORY_MODE=compress python memory_bot.py"],
           ["退出", "输入 quit"],
           ["生成文档", "python gen_docs.py（输出 docs/ 与桌面各一份）"]])

    h(doc, "六、文件清单", 1)
    table(doc,
          ["文件", "功能"],
          [["memory_bot.py", "记忆机器人主程序：三档会话记忆 + 按天持久化 + 关键词选取"],
           ["llm.py", "大模型调用封装 + 密钥加载（与第三讲共用）"],
           ["demo_run1.txt / demo_run2.txt", "课后任务演示记录（藏头诗→诗意追问→记忆验证）"],
           ["gen_docs.py", "生成本文档与《第四讲-讲解》"],
           ["requirements.txt", "依赖清单（jieba 可选，缺失自动降级）"],
           ["README.md", "仓库总说明与上手步骤"],
           ["docs/第四讲-讲解.docx", "课件逐节讲解 + 术语解释"],
           ["docs/Agent功能说明.docx", "本文档"]])

    save(doc, "Agent功能说明.docx")


if __name__ == "__main__":
    build_lecture_doc()
    build_agent_doc()
